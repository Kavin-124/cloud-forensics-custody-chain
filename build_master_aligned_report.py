import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from perfect_align_base import set_cell_margins, set_cell_background, set_table_borders
from perfect_chapters_content import inject_perfect_chapters

def make_page_break():
    p = OxmlElement('w:p')
    r = OxmlElement('w:r')
    br = OxmlElement('w:br')
    br.set(qn('w:type'), 'page')
    r.append(br)
    p.append(r)
    return p

def create_signature_table(doc):
    tbl = doc.add_table(rows=7, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, val="none")

    # Set column widths
    for row in tbl.rows:
        row.cells[0].width = Inches(3.2)
        row.cells[1].width = Inches(3.2)

    # Row 0: Signatures Header
    c0 = tbl.rows[0].cells[0]
    c0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    r0 = c0.paragraphs[0].add_run("SIGNATURE")
    r0.bold = True
    r0.font.name = "Times New Roman"
    r0.font.size = Pt(11)

    c1 = tbl.rows[0].cells[1]
    c1.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    r1 = c1.paragraphs[0].add_run("SIGNATURE")
    r1.bold = True
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(11)

    # Row 1: Space for physical signatures
    tbl.rows[1].cells[0].paragraphs[0].paragraph_format.space_before = Pt(28)
    tbl.rows[1].cells[1].paragraphs[0].paragraph_format.space_before = Pt(28)

    # Row 2: Names
    c0 = tbl.rows[2].cells[0]
    p = c0.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("Mr. K. MANIVANNAN, M.Tech., (Ph.D.)")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)

    c1 = tbl.rows[2].cells[1]
    p = c1.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("Mr. G.VIJAYKUMAR, M.E.,")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)

    # Row 3: Designations 1
    c0 = tbl.rows[3].cells[0]
    p = c0.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("HEAD OF THE DEPARTMENT")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)

    c1 = tbl.rows[3].cells[1]
    p = c1.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("SUPERVISOR")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)

    # Row 4: Designations 2
    c0 = tbl.rows[4].cells[0]
    p = c0.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("Department of Information Technology")
    r.font.name = "Times New Roman"
    r.font.size = Pt(10.5)

    c1 = tbl.rows[4].cells[1]
    p = c1.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("Assistant Professor")
    r.font.name = "Times New Roman"
    r.font.size = Pt(10.5)

    # Row 5: College Name
    c0 = tbl.rows[5].cells[0]
    p = c0.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("V.S.B. Engineering College")
    r.font.name = "Times New Roman"
    r.font.size = Pt(10.5)

    c1 = tbl.rows[5].cells[1]
    p = c1.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("Department of Information Technology")
    r.font.name = "Times New Roman"
    r.font.size = Pt(10.5)

    # Row 6: Location
    c0 = tbl.rows[6].cells[0]
    p = c0.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("Karur - 639 111.")
    r.font.name = "Times New Roman"
    r.font.size = Pt(10.5)

    c1 = tbl.rows[6].cells[1]
    p = c1.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("V.S.B. Engineering College, Karur - 639 111.")
    r.font.name = "Times New Roman"
    r.font.size = Pt(10.5)

    return tbl

def create_examiners_table(doc):
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, val="none")

    for row in tbl.rows:
        row.cells[0].width = Inches(3.2)
        row.cells[1].width = Inches(3.2)

    c0 = tbl.rows[0].cells[0]
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r0 = p0.add_run("INTERNAL EXAMINER")
    r0.bold = True
    r0.font.name = "Times New Roman"
    r0.font.size = Pt(11)

    c1 = tbl.rows[0].cells[1]
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = p1.add_run("EXTERNAL EXAMINER")
    r1.bold = True
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(11)

    return tbl

def build_perfect_aligned_report():
    src_path = r"C:\Users\kavin\.gemini\antigravity\brain\b37fd18a-0194-4357-93a6-9e420b486c27\.user_uploaded\media_1792547089640_5294a318.docx"
    doc = Document(src_path)

    # Locate the original Table of Contents table BEFORE any new tables are added
    toc_tbl = None
    for t in doc.tables:
        if len(t.rows) > 0 and len(t.rows[0].cells) > 0:
            if "CHAPTER NO." in t.rows[0].cells[0].text:
                toc_tbl = t
                break
    assert toc_tbl is not None, "Table of Contents table not found!"

    # =========================================================================
    # 1. COVER PAGE FORMATTING (P0 to P20)
    # =========================================================================
    p0 = doc.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after = Pt(4)
    if len(p0.runs) >= 4:
        p0.runs[2].text = ""
        p0.runs[3].text = "CRYPTOGRAPHIC CHAIN OF CUSTODY SYSTEM"
        p0.runs[3].font.name = "Times New Roman"
        p0.runs[3].font.size = Pt(17)
        p0.runs[3].bold = True

    p1 = doc.paragraphs[1]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(2)
    p1.paragraph_format.space_after = Pt(14)
    if len(p1.runs) >= 1:
        p1.runs[0].text = "FOR DIGITAL FORENSICS IN CLOUD ENVIRONMENTS"
        p1.runs[0].font.name = "Times New Roman"
        p1.runs[0].font.size = Pt(17)
        p1.runs[0].bold = True

    p2 = doc.paragraphs[2]
    p2.text = "A PROJECT REPORT"
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(6)
    p2.paragraph_format.space_after = Pt(14)
    if len(p2.runs) >= 1:
        p2.runs[0].font.name = "Times New Roman"
        p2.runs[0].font.size = Pt(13)
        p2.runs[0].bold = True

    # P4 has middle logos (TÜV and VSB emblem)
    p4 = doc.paragraphs[4]
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.paragraph_format.space_before = Pt(10)
    p4.paragraph_format.space_after = Pt(12)

    # P6: Submitted By
    p6 = doc.paragraphs[6]
    p6.text = "Submitted By"
    p6.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p6.paragraph_format.space_before = Pt(8)
    p6.paragraph_format.space_after = Pt(6)
    if len(p6.runs) >= 1:
        p6.runs[0].font.name = "Times New Roman"
        p6.runs[0].font.size = Pt(12)
        p6.runs[0].bold = True

    # Table 0: Student Names
    t0 = doc.tables[0]
    t0.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t0, val="none")
    t0.rows[0].cells[0].text = "KAVIN. R"
    t0.rows[0].cells[1].text = "(922524205075)"
    t0.rows[1].cells[0].text = "KESHAV. A. T"
    t0.rows[1].cells[1].text = "(922524205083)"

    for r_idx in range(2):
        for c_idx, cell in enumerate(t0.rows[r_idx].cells):
            set_cell_margins(cell, top=50, bottom=50, left=100, right=100)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.RIGHT
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(12)
                    r.bold = True

    # Degree Details (P12 to P20)
    degree_map = [
        (12, "in partial fulfillment for the award of the degree", 11, False, False, 14, 2),
        (13, "of", 11, False, False, 2, 2),
        (14, "BACHELOR OF TECHNOLOGY", 13, True, False, 2, 2),
        (15, "in", 11, False, False, 2, 2),
        (16, "INFORMATION TECHNOLOGY", 13, True, False, 2, 10),
        (18, "V.S.B. ENGINEERING COLLEGE, KARUR - 639 111 (Autonomous Institution)", 12, True, False, 8, 2),
        (19, "ANNA UNIVERSITY :: CHENNAI 600 025", 12, True, False, 2, 8),
        (20, "MAY 2026", 12, True, False, 4, 12)
    ]
    for p_idx, text_str, font_sz, is_bold, is_italic, sp_before, sp_after in degree_map:
        if p_idx < len(doc.paragraphs):
            p = doc.paragraphs[p_idx]
            p.text = text_str
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(sp_before)
            p.paragraph_format.space_after = Pt(sp_after)
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(font_sz)
                r.bold = is_bold
                r.italic = is_italic

    # Add page break after P20 (Cover page end)
    p20 = doc.paragraphs[20]
    p20._p.addnext(make_page_break())

    # =========================================================================
    # 2. BONAFIDE CERTIFICATE (P21 to P45)
    # =========================================================================
    p22 = doc.paragraphs[22]
    p22.text = "ANNA UNIVERSITY :: CHENNAI 600 025"
    p22.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p22.paragraph_format.space_before = Pt(16)
    p22.paragraph_format.space_after = Pt(4)
    for r in p22.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)
        r.bold = True

    p24 = doc.paragraphs[24]
    p24.text = "BONAFIDE CERTIFICATE"
    p24.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p24.paragraph_format.space_before = Pt(4)
    p24.paragraph_format.space_after = Pt(18)
    for r in p24.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(14)
        r.bold = True

    p25 = doc.paragraphs[25]
    p25.text = (
        "Certified that this project report titled “CRYPTOGRAPHIC CHAIN OF CUSTODY SYSTEM FOR DIGITAL FORENSICS "
        "IN CLOUD ENVIRONMENTS” is the bonafide work of “KAVIN. R (922524205075) and KESHAV. A. T (922524205083)” "
        "who carried out the project work under my supervision."
    )
    p25.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p25.paragraph_format.space_before = Pt(6)
    p25.paragraph_format.space_after = Pt(22)
    p25.paragraph_format.line_spacing = 1.5
    for r in p25.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)

    # Insert clean 2-column signature table right after P25
    sig_tbl = create_signature_table(doc)
    p25._p.addnext(sig_tbl._tbl)

    # Add Viva-Voce text and Examiners table after sig_tbl
    p_viva = doc.add_paragraph()
    p_viva.text = "Submitted for Anna University Project Viva-Voce held on ___________________"
    p_viva.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_viva.paragraph_format.space_before = Pt(36)
    p_viva.paragraph_format.space_after = Pt(28)
    for r in p_viva.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)
        r.italic = True
    sig_tbl._tbl.addnext(p_viva._p)

    exam_tbl = create_examiners_table(doc)
    p_viva._p.addnext(exam_tbl._tbl)

    # Page break after Bonafide Certificate
    pb_bonafide = make_page_break()
    exam_tbl._tbl.addnext(pb_bonafide)

    # Clear old messy paragraphs P26 through P45
    for p_i in range(26, 46):
        if p_i < len(doc.paragraphs):
            p_old = doc.paragraphs[p_i]
            p_old.text = ""
            p_old.paragraph_format.space_before = Pt(0)
            p_old.paragraph_format.space_after = Pt(0)

    # =========================================================================
    # 3. ACKNOWLEDGEMENT (P46 to P55)
    # =========================================================================
    p46 = doc.paragraphs[46]
    p46.text = "ACKNOWLEDGEMENT"
    p46.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p46.paragraph_format.space_before = Pt(16)
    p46.paragraph_format.space_after = Pt(14)
    for r in p46.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(14)
        r.bold = True

    ack_paragraphs = [
        "First and foremost, we express our thanks to our parents for providing us with a very supportive environment for completing this project. We wish to express our sincere thanks to our founder and Chairman Shri. V. S. BALSAMY B.Sc., L.L.B., for his endeavor in educating us in this premier institution.",
        "We extend our gratitude to Dr. C. VENNILA, M.E., Ph.D., Principal and Dr. T. S. KIRUBHASHANKAR, M.E., Ph.D., Vice Principal, V.S.B. ENGINEERING COLLEGE, KARUR for their high degree of encouragement and moral support during this project work.",
        "We express our indebtedness to Mr. K. MANIVANNAN, M.Tech., (Ph.D.), Head of the Department, Department of Information Technology for his continuous guidance throughout the course of our project.",
        "We are profoundly grateful to our Project Supervisor Mr. G.VIJAYKUMAR, M.E., Assistant Professor, Department of Information Technology for his valuable technical support, critical reviews, and encouragement.",
        "Our sincere thanks to all the faculty members of the Department of Information Technology, V.S.B. Engineering College and our friends for their help in the successful completion of this project work.",
        "Finally, we bow before God, the almighty who always had a better plan for us. We give our praise and glory to almighty God for the successful completion of our project."
    ]
    for idx, ack_txt in enumerate(ack_paragraphs):
        p_target = doc.paragraphs[48 + idx]
        p_target.text = ack_txt
        p_target.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_target.paragraph_format.space_before = Pt(0)
        p_target.paragraph_format.space_after = Pt(6)
        p_target.paragraph_format.line_spacing = 1.5
        for r in p_target.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)

    # Page break after Acknowledgement
    p53 = doc.paragraphs[53]
    p53._p.addnext(make_page_break())

    # =========================================================================
    # 4. DEPARTMENT VISION & MISSION, PEOs, POs, PSOs (P56 to P95)
    # =========================================================================
    p56 = doc.paragraphs[56]
    p56.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p56.paragraph_format.space_before = Pt(12)
    p56.paragraph_format.space_after = Pt(10)

    p57 = doc.paragraphs[57]
    p57.text = "DEPARTMENT OF INFORMATION TECHNOLOGY"
    p57.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p57.paragraph_format.space_before = Pt(6)
    p57.paragraph_format.space_after = Pt(12)
    for r in p57.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(14)
        r.bold = True

    # Format vision, mission, peo, po, pso paragraphs
    for p_i in range(58, 96):
        if p_i < len(doc.paragraphs):
            p = doc.paragraphs[p_i]
            txt = p.text.strip()
            if not txt:
                continue
            if txt.startswith(("Vision", "Mission", "Program")):
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.keep_with_next = True
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(12)
                    r.bold = True
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.5
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(12)

    # Page break after PSO3 (P94)
    p94 = doc.paragraphs[94]
    p94._p.addnext(make_page_break())

    # =========================================================================
    # 5. ABSTRACT (P96 to P102)
    # =========================================================================
    p96 = doc.paragraphs[96]
    p96.text = "ABSTRACT"
    p96.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p96.paragraph_format.space_before = Pt(18)
    p96.paragraph_format.space_after = Pt(14)
    for r in p96.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(14)
        r.bold = True

    for p_idx in [98, 99, 100, 101]:
        if p_idx < len(doc.paragraphs):
            p = doc.paragraphs[p_idx]
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.5
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(12)

    # Page break after Abstract
    p101 = doc.paragraphs[101]
    p101._p.addnext(make_page_break())

    # =========================================================================
    # 6. TABLE OF CONTENTS (P103 and Table 1)
    # =========================================================================
    p103 = doc.paragraphs[103]
    p103.text = "TABLE OF CONTENTS"
    p103.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p103.paragraph_format.space_before = Pt(18)
    p103.paragraph_format.space_after = Pt(14)
    for r in p103.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(14)
        r.bold = True

    t1 = toc_tbl
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t1, color="B0C4DE", sz="6", val="single")

    # Header Row
    hdr_row = t1.rows[0]
    for c_idx, cell in enumerate(hdr_row.cells):
        set_cell_background(cell, "1F497D")
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(10)
            r.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

    # Data Rows
    for r_idx in range(1, len(t1.rows)):
        row = t1.rows[r_idx]
        bg_col = "F2F5F9" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell in enumerate(row.cells):
            if bg_col != "FFFFFF":
                set_cell_background(cell, bg_col)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 1 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.15
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(9.5)

    # =========================================================================
    # 7. REMOVE OLD PARAGRAPHS AFTER TABLE 1 & INJECT PERFECT CHAPTERS
    # =========================================================================
    body = doc._body._body
    t1_elem = t1._tbl
    t1_idx = body.index(t1_elem)

    # Remove all elements after Table 1 except the final sectPr element
    elements_to_remove = [child for child in body[t1_idx + 1:-1]]
    for el in elements_to_remove:
        body.remove(el)
    print(f"Removed {len(elements_to_remove)} old paragraphs after Table 1.")

    # Page break after TOC
    doc.add_page_break()

    # Now inject the expansive, meticulously aligned chapters!
    inject_perfect_chapters(doc)

    out_file = "VSB_IT_Cryptographic_Chain_of_Custody_Perfect_Aligned.docx"
    doc.save(out_file)
    print(f"Successfully generated and saved '{out_file}'!")
    return out_file

if __name__ == "__main__":
    build_perfect_aligned_report()
